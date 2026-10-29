"""Actual Word packages, owned images, formulas and protected downloads."""
import io
import zipfile

import pytest
from docx import Document
from lxml import etree

from app.api.routes import ai
from app.services.lesson_word import build_lesson_docx, WORD_MIME
from app.services.lesson_materials import asset_url
from app.services.lesson_content import clean_lesson_content
from test_textbook_pipeline import environment, import_reviewed, generation_payload


SOURCE = r'''# 牛顿第一定律教案

## 教学目标

理解**惯性**，观察实验并表达结论。

1. 观察小球运动。
2. 讨论物体运动条件。

| 教学环节 | 教师活动 | 学生活动 | 时间 |
| --- | --- | --- | --- |
| 新授 | 组织观察，提示比较条件。 | 记录观察结果并交流。 | 5分钟 |

加速度公式为 $a=\frac{F}{m}$，平方关系为 \(s=t^2\)。

$$v=\sqrt{2as}$$
'''


def test_word_package_has_real_tables_math_styles_and_readable_chinese():
    raw=build_lesson_docx(SOURCE,'牛顿第一定律')
    assert raw.startswith(b'PK')
    document=Document(io.BytesIO(raw))
    assert len(document.tables)==1
    assert len(document.tables[0].rows)==2
    text='\n'.join(p.text for p in document.paragraphs)
    assert '牛顿第一定律教案' in text and '理解惯性' in text
    assert text.count('牛顿第一定律教案')==1
    assert '#' not in text and '**' not in text and r'\frac' not in text
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        xml=etree.fromstring(archive.read('word/document.xml'))
        assert xml.xpath('count(//m:f)',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'})==1
        assert xml.xpath('count(//m:rad)',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'})==1
        assert b'D9D9D9' in archive.read('word/document.xml')
        assert b'Title' in archive.read('word/document.xml')
        styles=etree.fromstring(archive.read('word/styles.xml'))
        namespace={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        for name in ('Title','Heading1','Heading2','Heading3','TitleChar','Heading1Char'):
            style=styles.xpath(f'//w:style[@w:styleId="{name}"]',namespaces=namespace)[0]
            assert not style.xpath('.//w:pBdr|.//w:color[@w:themeColor]',namespaces=namespace)


def test_preparation_omits_figures_and_source_section_keeps_teaching_process():
    image='![实验图](/api/v1/textbooks/assets/tb-'+ '1'*32 +'/image)'
    source='''# 运动教案
## 三、教学准备
| 类别 | 内容 |
| --- | --- |
| 多媒体 | 投影仪、电脑、教材课件（包含图2.4-1、图2.4-4等） |
| 材料 | 米尺、铁球 |
> ⚠️ **特别说明**：
> - 图2.4-1需展示对比照片，突出空气阻力。
'''+image+'''
[[图片：牛顿管实验]]
### 教师准备
准备铁球和纸张。
## 四、教学过程
展示图2.4-1，讨论实验条件。
'''+image+'''
## 五、教材依据
测试教材 PDF第1–2页。
### 来源说明
仅作为内部依据。
## 六、教学反思
关注学生的实验解释。
'''
    cleaned=clean_lesson_content(source)
    prep=cleaned.split('## 三、教学准备')[1].split('## 四、教学过程')[0]
    assert '投影仪' in prep and '米尺' in prep and '准备铁球和纸张' in prep
    assert '图2.4' not in prep and '![' not in prep and '[[图片' not in prep and '特别说明' not in prep
    assert cleaned.count(image)==1
    assert '教材依据' not in cleaned and 'PDF第' not in cleaned and '来源说明' not in cleaned
    assert '教学反思' in cleaned and '讨论实验条件' in cleaned


def test_spaced_formulas_and_table_images_stay_inside_cell(tmp_path):
    from PIL import Image
    identifier='tb-'+'2'*32
    path=tmp_path/'figure.png';Image.new('RGB',(1200,500),'white').save(path)
    source='加速度为 $ a = \\frac{F}{m} $。\n\n| 环节 | 教学活动 | 设计意图 | 时间 |\n|---|---|---|---|\n|新授|![实验](/api/v1/textbooks/assets/'+identifier+'/image)|观察并解释|5分钟|'
    raw=build_lesson_docx(source,'教案',{identifier:path})
    document=Document(io.BytesIO(raw))
    cell=document.tables[0].cell(1,1)
    assert document.inline_shapes[0].width < cell.width
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        xml=etree.fromstring(archive.read('word/document.xml'))
        assert xml.xpath('count(//m:f)',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'})==1
        assert '\\frac' not in ''.join(document.paragraphs[1].text)


def test_tall_figures_have_bounded_height_and_stay_with_introduction(tmp_path):
    from PIL import Image
    from docx.shared import Cm
    identifier='tb-'+'3'*32
    path=tmp_path/'tall.png';Image.new('RGB',(500,1400),'white').save(path)
    source='展示实验图：\n\n![图2.4-1 图2.4-1 实验](/api/v1/textbooks/assets/'+identifier+'/image)\n\n。\n\n观察实验。'
    document=Document(io.BytesIO(build_lesson_docx(source,'实验',{identifier:path})))
    assert document.inline_shapes[0].height<=Cm(10)+100
    introduction=next(paragraph for paragraph in document.paragraphs if paragraph.text=='展示实验图：')
    assert introduction.paragraph_format.keep_with_next
    text='\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert '图2.4-1 图2.4-1' not in text and '\n。\n' not in text


@pytest.mark.asyncio
async def test_generation_returns_word_with_embedded_image_and_owner_permissions(environment,monkeypatch):
    client,sessions,root=environment
    book,chapter,image,other=await import_reviewed(client)
    class Model:
        async def generate_lesson_plan(self,**kwargs):
            return SOURCE+'\n\n![伽利略实验]('+asset_url(image['asset_id'])+')'
    monkeypatch.setattr(ai,'get_dynamic_ai_service',Model)
    response=await client.post('/ai/generate',json=generation_payload(chapter['id']))
    assert response.status_code==200,response.text
    generated=response.json()['data']
    assert generated['document_format']=='docx'
    assert generated['document_filename'].endswith('.docx')
    endpoint='/ai/lesson-documents/'+generated['document_id']
    response=await client.get(endpoint)
    assert response.status_code==200
    assert response.headers['content-type']==WORD_MIME
    assert '.docx' in response.headers['content-disposition']
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        media=[name for name in archive.namelist() if name.startswith('word/media/')]
        assert len(media)==1
        assert archive.read(media[0])==(await client.get(asset_url(image['asset_id']).removeprefix('/api/v1'))).content
        assert b'http://example' not in archive.read('word/_rels/document.xml.rels')
    assert (await client.get(endpoint,headers={'x-user':'2'})).status_code==404
    assert (await client.get(endpoint,headers={'x-user':''})).status_code==401
    assert (await client.post('/ai/lesson-documents',headers={'x-user':'2'},json={'lesson_id':generated['lesson_id']})).status_code==404
    # An owned historical lesson can become Word without running the model again.
    converted=await client.post('/ai/lesson-documents',json={'lesson_id':generated['lesson_id'],'title':'历史教案'})
    assert converted.status_code==200,converted.text


@pytest.mark.asyncio
async def test_legacy_general_source_conversion_never_fetches_arbitrary_images(environment):
    client,_,_=environment
    response=await client.post('/ai/lesson-documents',json={'title':'旧教案','content':SOURCE+'\n\n![外部图片](https://example.invalid/secret.png)\n\n![本地文件](file:///C:/private.png)'})
    assert response.status_code==200,response.text
    document=await client.get('/ai/lesson-documents/'+response.json()['data']['document_id'])
    with zipfile.ZipFile(io.BytesIO(document.content)) as archive:
        assert not any(name.startswith('word/media/') for name in archive.namelist())
        assert '图片待补充'.encode() in archive.read('word/document.xml')
        assert b'example.invalid' not in archive.read('word/_rels/document.xml.rels')
    assert (await client.post('/ai/lesson-documents',json={'content':' '})).status_code==422


@pytest.mark.asyncio
async def test_existing_word_link_is_rebuilt_with_current_rules(environment):
    from app.services.lesson_documents import DOCUMENT_VERSION
    from app.services.textbook_service import private_path
    import json
    client,_,_=environment
    response=await client.post('/ai/lesson-documents',json={'title':'旧教案','content':SOURCE+'\n\n## 教材依据\n测试教材来源。'})
    identifier=response.json()['data']['document_id']
    metadata_path=private_path(f'word-documents/1/{identifier}.json')
    metadata=json.loads(metadata_path.read_text(encoding='utf-8'));metadata['version']=0
    metadata_path.write_text(json.dumps(metadata),encoding='utf-8')
    private_path(f'word-documents/1/{identifier}.docx').write_bytes(b'old-format-sentinel')
    response=await client.get('/ai/lesson-documents/'+identifier)
    assert response.status_code==200 and response.content.startswith(b'PK')
    assert json.loads(metadata_path.read_text(encoding='utf-8'))['version']==DOCUMENT_VERSION
    document=Document(io.BytesIO(response.content))
    text='\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert '教材依据' not in text and '测试教材来源' not in text
