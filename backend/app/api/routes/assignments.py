"""
Assignment API routes.
Handles student-instructor group assignments.
"""
from typing import List
from fastapi import APIRouter, Depends, HTTPException, status, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy import select, func, delete
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.database import get_db
from app.core.deps import require_admin, get_current_user
from app.models.user import User, UserRole
from app.models.semester import Semester
from app.models.group import Group
from app.schemas.group import (
    GroupCreate,
    GroupResponse,
    BatchAssignment,
    BatchUnassignment,
    InstructorWithStudents,
    AssignmentOverview,
    AssignmentImportResult,
    AutoAssignRequest,
)
from app.schemas.user import UserResponse
from app.utils.excel_handler import parse_excel_file, create_assignment_import_template

router = APIRouter()


@router.get(
    "/overview/{semester_id}",
    response_model=AssignmentOverview,
    summary="获取分配概览",
)
async def get_assignment_overview(
    semester_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
) -> AssignmentOverview:
    """
    获取指定批次的学生分配概览。
    
    返回：
    - 指导老师总数
    - 学生总数
    - 已分配/未分配学生数
    - 每个老师名下的学生列表（用于拖拽分配界面）
    """
    # 验证批次存在
    semester = await db.get(Semester, semester_id)
    if not semester:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="批次不存在",
        )
    
    # 获取所有指导老师
    instructors_result = await db.execute(
        select(User).where(
            User.role == UserRole.INSTRUCTOR,
            User.is_active == True,
        )
    )
    instructors = instructors_result.scalars().all()
    
    # 获取所有学生
    students_result = await db.execute(
        select(User).where(
            User.role == UserRole.STUDENT,
            User.is_active == True,
        )
    )
    all_students = students_result.scalars().all()
    
    # 获取该批次的所有分组
    groups_result = await db.execute(
        select(Group)
        .where(Group.semester_id == semester_id)
        .options(
            selectinload(Group.instructor),
            selectinload(Group.student),
        )
    )
    groups = groups_result.scalars().all()
    
    # 构建每个老师的学生列表
    instructor_students_map: dict[int, list[User]] = {i.id: [] for i in instructors}
    assigned_student_ids: set[int] = set()
    
    for group in groups:
        if group.instructor_id in instructor_students_map:
            instructor_students_map[group.instructor_id].append(group.student)
            assigned_student_ids.add(group.student_id)
    
    # 构建返回数据
    groups_data = []
    for instructor in instructors:
        students = instructor_students_map.get(instructor.id, [])
        groups_data.append(
            InstructorWithStudents(
                instructor=UserResponse.model_validate(instructor),
                students=[UserResponse.model_validate(s) for s in students],
                student_count=len(students),
            )
        )
    
    return AssignmentOverview(
        semester_id=semester_id,
        total_instructors=len(instructors),
        total_students=len(all_students),
        assigned_students=len(assigned_student_ids),
        unassigned_students=len(all_students) - len(assigned_student_ids),
        groups=groups_data,
    )


@router.get(
    "/import/template",
    summary="下载分配导入模板",
)
async def download_assignment_import_template(
    _: User = Depends(require_admin),
) -> StreamingResponse:
    """下载学生-指导老师分配的Excel导入模板。"""
    return create_assignment_import_template()


@router.post(
    "/import",
    response_model=AssignmentImportResult,
    summary="批量导入学生分配",
)
async def import_assignments(
    semester_id: int = Query(..., description="批次ID"),
    file: UploadFile = File(..., description="Excel文件"),
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
) -> AssignmentImportResult:
    """通过Excel批量导入学生与指导老师的分配关系。"""
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="请上传Excel文件 (.xlsx 或 .xls)",
        )

    content = await file.read()

    try:
        rows = parse_excel_file(content)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Excel文件读取失败: {str(e)}",
        )

    semester = await db.get(Semester, semester_id)
    if not semester:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="批次不存在",
        )

    success_count = 0
    fail_count = 0
    failed_items = []

    for index, row in enumerate(rows, start=2):
        student_code = str(row.get("学生学号") or "").strip()
        instructor_code = str(row.get("指导老师工号") or "").strip()

        if not student_code or not instructor_code:
            failed_items.append({
                "row": index,
                "data": row,
                "reason": "必填字段缺失",
            })
            fail_count += 1
            continue

        try:
            # 查找学生
            student_result = await db.execute(
                select(User).where(
                    User.user_code == student_code,
                    User.role == UserRole.STUDENT,
                    User.is_active == True,
                )
            )
            student = student_result.scalar_one_or_none()
            if not student:
                failed_items.append({
                    "row": index,
                    "data": row,
                    "reason": f"学生不存在或角色不正确: {student_code}",
                })
                fail_count += 1
                continue

            # 查找指导老师
            instructor_result = await db.execute(
                select(User).where(
                    User.user_code == instructor_code,
                    User.role == UserRole.INSTRUCTOR,
                    User.is_active == True,
                )
            )
            instructor = instructor_result.scalar_one_or_none()
            if not instructor:
                failed_items.append({
                    "row": index,
                    "data": row,
                    "reason": f"指导老师不存在或角色不正确: {instructor_code}",
                })
                fail_count += 1
                continue

            # 更新或创建分配关系
            group_result = await db.execute(
                select(Group).where(
                    Group.semester_id == semester_id,
                    Group.student_id == student.id,
                )
            )
            group = group_result.scalar_one_or_none()

            if group:
                group.instructor_id = instructor.id
            else:
                db.add(
                    Group(
                        semester_id=semester_id,
                        instructor_id=instructor.id,
                        student_id=student.id,
                    )
                )

            success_count += 1

        except Exception as e:  # pragma: no cover - defensive
            failed_items.append({
                "row": index,
                "data": row,
                "reason": str(e),
            })
            fail_count += 1

    await db.commit()

    return AssignmentImportResult(
        success_count=success_count,
        fail_count=fail_count,
        failed_items=failed_items,
    )


@router.get(
    "/unassigned/{semester_id}",
    response_model=List[UserResponse],
    summary="获取未分配学生列表",
)
async def get_unassigned_students(
    semester_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
) -> List[UserResponse]:
    """获取指定批次中尚未分配给任何老师的学生列表。"""
    # 获取已分配的学生ID
    assigned_result = await db.execute(
        select(Group.student_id).where(Group.semester_id == semester_id)
    )
    assigned_ids = {row[0] for row in assigned_result.fetchall()}
    
    # 获取所有活跃学生中未分配的
    query = select(User).where(
        User.role == UserRole.STUDENT,
        User.is_active == True,
    )
    if assigned_ids:
        query = query.where(User.id.notin_(assigned_ids))
    
    result = await db.execute(query)
    students = result.scalars().all()
    
    return [UserResponse.model_validate(s) for s in students]


@router.post(
    "/batch",
    response_model=dict,
    summary="批量分配学生给老师",
)
async def batch_assign_students(
    assignment: BatchAssignment,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
) -> dict:
    """
    批量将学生分配给指定的指导老师。
    
    用于拖拽分配或批量勾选场景。
    
    注意：一个学生在一个批次内只能绑定一个指导老师。
    如果学生已被分配给其他老师，会先解除原有分配。
    """
    # 验证批次存在
    semester = await db.get(Semester, assignment.semester_id)
    if not semester:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="批次不存在",
        )
    
    # 验证老师存在且角色正确
    instructor = await db.get(User, assignment.instructor_id)
    if not instructor or instructor.role != UserRole.INSTRUCTOR:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="指导老师不存在或角色不正确",
        )
    
    # 验证所有学生存在且角色正确
    students_result = await db.execute(
        select(User).where(
            User.id.in_(assignment.student_ids),
            User.role == UserRole.STUDENT,
            User.is_active == True,
        )
    )
    valid_students = students_result.scalars().all()
    valid_student_ids = {s.id for s in valid_students}
    
    if len(valid_student_ids) != len(assignment.student_ids):
        invalid_ids = set(assignment.student_ids) - valid_student_ids
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"以下学生ID无效: {invalid_ids}",
        )
    
    # 删除这些学生在该批次的现有分配
    await db.execute(
        delete(Group).where(
            Group.semester_id == assignment.semester_id,
            Group.student_id.in_(assignment.student_ids),
        )
    )
    
    # 创建新的分配关系
    new_groups = [
        Group(
            semester_id=assignment.semester_id,
            instructor_id=assignment.instructor_id,
            student_id=student_id,
        )
        for student_id in assignment.student_ids
    ]
    db.add_all(new_groups)
    await db.commit()
    
    return {
        "message": f"成功将 {len(new_groups)} 名学生分配给老师",
        "assigned_count": len(new_groups),
    }


@router.post(
    "/single",
    response_model=GroupResponse,
    summary="单个学生分配",
)
async def assign_single_student(
    group_data: GroupCreate,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
) -> GroupResponse:
    """
    将单个学生分配给指定的指导老师。
    
    如果该学生在该批次已有分配，会更新为新的老师。
    """
    # 验证批次、老师、学生
    semester = await db.get(Semester, group_data.semester_id)
    if not semester:
        raise HTTPException(status_code=404, detail="批次不存在")
    
    instructor = await db.get(User, group_data.instructor_id)
    if not instructor or instructor.role != UserRole.INSTRUCTOR:
        raise HTTPException(status_code=400, detail="指导老师不存在或角色不正确")
    
    student = await db.get(User, group_data.student_id)
    if not student or student.role != UserRole.STUDENT:
        raise HTTPException(status_code=400, detail="学生不存在或角色不正确")
    
    # 检查是否已存在分配，存在则更新
    existing_result = await db.execute(
        select(Group).where(
            Group.semester_id == group_data.semester_id,
            Group.student_id == group_data.student_id,
        )
    )
    existing_group = existing_result.scalar_one_or_none()
    
    if existing_group:
        existing_group.instructor_id = group_data.instructor_id
        group = existing_group
    else:
        group = Group(
            semester_id=group_data.semester_id,
            instructor_id=group_data.instructor_id,
            student_id=group_data.student_id,
        )
        db.add(group)
    
    await db.commit()
    await db.refresh(group)
    
    return GroupResponse.model_validate(group)


@router.delete(
    "/{semester_id}/{student_id}",
    summary="取消单个学生的分配",
)
async def unassign_student(
    semester_id: int,
    student_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
) -> dict:
    """取消指定学生在指定批次的分配关系。"""
    # 查找分组
    result = await db.execute(
        select(Group).where(
            Group.semester_id == semester_id,
            Group.student_id == student_id,
        )
    )
    group = result.scalar_one_or_none()
    
    if not group:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="未找到该学生的分配记录"
        )
    
    await db.delete(group)
    await db.commit()
    
    return {"message": "取消分配成功"}


@router.delete(
    "/batch",
    summary="批量取消分配",
)
async def batch_unassign_students(
    unassignment: BatchUnassignment,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
) -> dict:
    """批量取消学生的分配关系。"""
    result = await db.execute(
        delete(Group).where(
            Group.semester_id == unassignment.semester_id,
            Group.student_id.in_(unassignment.student_ids),
        )
    )
    await db.commit()
    
    return {
        "message": f"成功取消 {result.rowcount} 名学生的分配",
        "removed_count": result.rowcount,
    }


@router.get(
    "/instructor/{semester_id}",
    response_model=List[UserResponse],
    summary="获取老师名下的学生",
)
async def get_instructor_students(
    semester_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> List[UserResponse]:
    """
    获取当前老师在指定批次名下的学生列表。
    
    指导老师专用接口。
    """
    if current_user.role not in [UserRole.INSTRUCTOR, UserRole.ADMIN]:
        raise HTTPException(status_code=403, detail="权限不足")
    
    instructor_id = current_user.id
    
    # 获取分组
    groups_result = await db.execute(
        select(Group)
        .where(
            Group.semester_id == semester_id,
            Group.instructor_id == instructor_id,
        )
        .options(selectinload(Group.student))
    )
    groups = groups_result.scalars().all()
    
    return [UserResponse.model_validate(g.student) for g in groups]


@router.post(
    "/auto-assign/{semester_id}",
    summary="一键自动分配",
)
async def auto_assign_students(
    semester_id: int,
    request: AutoAssignRequest,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
) -> dict:
    """
    一键自动分配学生给指导老师。
    
    算法：将学生尽可能平均分配给指定的指导老师。
    例如：52个学生分给8个老师，会有4个老师分到7人，4个老师分到6人。
    """
    # 验证批次存在
    semester = await db.get(Semester, semester_id)
    if not semester:
        raise HTTPException(status_code=404, detail="批次不存在")
    
    # 获取已分配的学生ID
    assigned_result = await db.execute(
        select(Group.student_id).where(Group.semester_id == semester_id)
    )
    assigned_ids = {row[0] for row in assigned_result.fetchall()}
    
    # 获取未分配学生（可按条件筛选）
    query = select(User).where(
        User.role == UserRole.STUDENT,
        User.is_active == True,
    )
    if assigned_ids:
        query = query.where(User.id.notin_(assigned_ids))
    if request.class_name:
        query = query.where(User.class_name == request.class_name)
    if request.college:
        query = query.where(User.college == request.college)
    if request.major:
        query = query.where(User.major == request.major)
    
    # 按班级排序，确保同班学生在一起
    query = query.order_by(User.class_name, User.user_code)
    
    result = await db.execute(query)
    students = list(result.scalars().all())
    
    if not students:
        return {"message": "没有需要分配的学生", "assigned_count": 0}
    
    # 获取指定的指导老师（或全部老师）
    instructor_query = select(User).where(
        User.role == UserRole.INSTRUCTOR,
        User.is_active == True,
    )
    if request.instructor_ids:
        instructor_query = instructor_query.where(User.id.in_(request.instructor_ids))
    instructor_query = instructor_query.order_by(User.id)
    
    instructors_result = await db.execute(instructor_query)
    instructors = list(instructors_result.scalars().all())
    
    if not instructors:
        raise HTTPException(status_code=400, detail="没有可用的指导老师")
    
    # 获取每个老师当前已分配的学生数量
    instructor_counts = {i.id: 0 for i in instructors}
    for group in (await db.execute(
        select(Group).where(Group.semester_id == semester_id)
    )).scalars().all():
        if group.instructor_id in instructor_counts:
            instructor_counts[group.instructor_id] += 1
    
    new_groups = []
    
    # 如果启用同班级分组
    if request.same_class_grouping:
        # 按班级分组学生
        from collections import defaultdict
        class_groups = defaultdict(list)
        for student in students:
            class_name = student.class_name or "未分班"
            class_groups[class_name].append(student)
        
        # 计算总学生数和每个班级的学生数
        total_students = len(students)
        total_instructors = len(instructors)
        
        # 按班级学生数比例分配老师
        class_instructor_allocation = {}
        allocated_instructors = 0
        
        # 将班级按人数排序（大班级优先），确保大班级优先分配
        sorted_classes = sorted(
            class_groups.items(), 
            key=lambda x: len(x[1]), 
            reverse=True
        )
        
        # 第一步：按比例计算每个班级应该分配多少老师
        for idx, (class_name, class_students) in enumerate(sorted_classes):
            class_size = len(class_students)
            # 按学生比例计算该班级应分配的老师数（向下取整）
            if idx == len(sorted_classes) - 1:
                # 最后一个班级分配剩余的所有老师
                instructors_for_class = total_instructors - allocated_instructors
            else:
                # 按比例分配，至少分配1个老师
                instructors_for_class = max(1, int(class_size / total_students * total_instructors))
            
            class_instructor_allocation[class_name] = instructors_for_class
            allocated_instructors += instructors_for_class
        
        # 第二步：为每个班级分配具体的老师
        instructor_pool = list(instructors)
        instructor_index = 0
        
        for class_name, class_students in sorted_classes:
            # 获取该班级应分配的老师数量
            num_instructors = class_instructor_allocation[class_name]
            
            # 从老师池中取出相应数量的老师
            class_instructors = []
            for _ in range(num_instructors):
                if instructor_index < len(instructor_pool):
                    class_instructors.append(instructor_pool[instructor_index].id)
                    instructor_index += 1
            
            # 如果没有老师可分配（理论上不应该发生）
            if not class_instructors:
                raise HTTPException(
                    status_code=400,
                    detail=f"班级 {class_name} 无法分配老师"
                )
            
            # 第三步：在班级内部，将学生平均分配给这些老师
            for student_idx, student in enumerate(class_students):
                # 轮流分配给该班级的老师
                instructor_id = class_instructors[student_idx % len(class_instructors)]
                
                new_groups.append(Group(
                    semester_id=semester_id,
                    instructor_id=instructor_id,
                    student_id=student.id,
                ))
                instructor_counts[instructor_id] += 1
    else:
        # 平均分配算法：每次将学生分配给当前学生数最少的老师
        for student in students:
            # 找到当前学生数最少的老师
            min_instructor_id = min(instructor_counts, key=instructor_counts.get)
            
            new_groups.append(Group(
                semester_id=semester_id,
                instructor_id=min_instructor_id,
                student_id=student.id,
            ))
            instructor_counts[min_instructor_id] += 1
    
    db.add_all(new_groups)
    await db.commit()
    
    return {
        "message": f"成功自动分配 {len(new_groups)} 名学生给 {len(instructors)} 位老师",
        "assigned_count": len(new_groups),
    }


@router.get(
    "/filter/options/{semester_id}",
    summary="获取筛选选项",
)
async def get_filter_options(
    semester_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
) -> dict:
    """获取未分配学生的学院、专业、班级筛选选项。"""
    # 获取已分配的学生ID
    assigned_result = await db.execute(
        select(Group.student_id).where(Group.semester_id == semester_id)
    )
    assigned_ids = {row[0] for row in assigned_result.fetchall()}
    
    # 获取未分配学生
    query = select(User).where(
        User.role == UserRole.STUDENT,
        User.is_active == True,
    )
    if assigned_ids:
        query = query.where(User.id.notin_(assigned_ids))
    
    result = await db.execute(query)
    students = result.scalars().all()
    
    colleges = set()
    majors = set()
    classes = set()
    
    for s in students:
        if s.college:
            colleges.add(s.college)
        if s.major:
            majors.add(s.major)
        if s.class_name:
            classes.add(s.class_name)
    
    return {
        "colleges": sorted(list(colleges)),
        "majors": sorted(list(majors)),
        "classes": sorted(list(classes)),
    }


@router.get(
    "/student/my-instructor/{semester_id}",
    response_model=UserResponse,
    summary="获取学生的指导老师",
)
async def get_my_instructor(
    semester_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> UserResponse:
    """
    获取当前学生在指定批次的指导老师信息。
    
    学生专用接口。
    """
    if current_user.role != UserRole.STUDENT:
        raise HTTPException(status_code=403, detail="仅学生可访问")
    
    group_result = await db.execute(
        select(Group)
        .where(
            Group.semester_id == semester_id,
            Group.student_id == current_user.id,
        )
        .options(selectinload(Group.instructor))
    )
    group = group_result.scalar_one_or_none()
    
    if not group:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="您在该批次尚未被分配指导老师",
        )
    
    return UserResponse.model_validate(group.instructor)


@router.get(
    "/export/{semester_id}",
    summary="导出分组分配情况",
)
async def export_assignments(
    semester_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
) -> StreamingResponse:
    """导出指定批次的分组分配情况为Excel文件。"""
    import io
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    
    # 验证批次存在
    semester = await db.get(Semester, semester_id)
    if not semester:
        raise HTTPException(status_code=404, detail="批次不存在")
    
    # 获取分组数据
    result = await db.execute(
        select(Group)
        .where(Group.semester_id == semester_id)
        .options(selectinload(Group.instructor), selectinload(Group.student))
        .order_by(Group.instructor_id, Group.student_id)
    )
    groups = result.scalars().all()
    
    # 创建工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "分组分配情况"
    
    # 设置表头
    headers = ["序号", "指导老师工号", "指导老师姓名", "学生学号", "学生姓名", 
               "学生学院", "学生专业", "学生班级", "分配时间"]
    ws.append(headers)
    
    # 设置表头样式
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 写入数据
    for idx, group in enumerate(groups, start=1):
        ws.append([
            idx,
            group.instructor.user_code if group.instructor else "",
            group.instructor.name if group.instructor else "",
            group.student.user_code if group.student else "",
            group.student.name if group.student else "",
            group.student.college or "",
            group.student.major or "",
            group.student.class_name or "",
            group.created_at.strftime("%Y-%m-%d %H:%M:%S") if group.created_at else "",
        ])
    
    # 调整列宽
    column_widths = [8, 15, 12, 15, 12, 20, 20, 15, 20]
    for idx, width in enumerate(column_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width
    
    # 保存到字节流
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"分组分配_{semester.name}_{semester.year}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
    )
