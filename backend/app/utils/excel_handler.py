"""
Excel file handling utilities.
"""
import io
from typing import List, Dict, Any
from fastapi import UploadFile
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def parse_excel_file(file_content: bytes) -> List[Dict[str, Any]]:
    """
    解析Excel文件内容。
    
    Args:
        file_content: Excel文件的字节内容
    
    Returns:
        解析后的数据列表，每行为一个字典
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_content))
    ws = wb.active
    
    # 获取表头
    headers = []
    for cell in ws[1]:
        if cell.value:
            headers.append(str(cell.value).strip())
    
    # 解析数据行
    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):  # 跳过空行
            continue
        row_data = {}
        for i, value in enumerate(row):
            if i < len(headers):
                row_data[headers[i]] = value
        data.append(row_data)
    
    return data


def create_excel_response(
    data: List[Dict[str, Any]],
    filename: str = "export.xlsx",
    sheet_name: str = "Sheet1",
) -> StreamingResponse:
    """
    创建Excel文件下载响应。
    
    Args:
        data: 要导出的数据列表
        filename: 下载文件名
        sheet_name: 工作表名称
    
    Returns:
        FastAPI StreamingResponse
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    if not data:
        # 空数据返回空表
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )
    
    # 获取表头
    headers = list(data[0].keys())
    
    # 样式定义
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    
    # 写入表头
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # 写入数据
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, header in enumerate(headers, 1):
            value = row_data.get(header)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 调整列宽
    for col_idx, header in enumerate(headers, 1):
        max_length = len(str(header))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max_length + 2
    
    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def create_user_import_template() -> StreamingResponse:
    """
    创建用户导入模板Excel文件。
    
    Returns:
        FastAPI StreamingResponse
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "用户导入模板"
    
    # 表头
    headers = [
        "工号/学号",
        "姓名",
        "角色",
        "学院",
        "专业",
        "班级",
        "班主任",
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    # 示例数据
    examples = [
        ["T001", "张教授", "instructor", "教育学院", "教育学", "教学一班", "王老师"],
        ["S2021001", "李小明", "student", "教育学院", "小学教育", "2021级1班", "张班主任"],
    ]
    
    for row_idx, example in enumerate(examples, 2):
        for col_idx, value in enumerate(example, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # 添加说明
    ws.cell(row=5, column=1, value="说明：")
    ws.cell(row=6, column=1, value="角色可选值: admin(管理员), instructor(指导老师), student(学生)")
    ws.cell(row=7, column=1, value="默认密码为工号/学号")
    
    # 调整列宽
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 16
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=user_import_template.xlsx"},
    )


def create_assignment_import_template() -> StreamingResponse:
    """创建学生分配导入模板Excel文件。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "分配导入模板"
    
    # 表头
    headers = [
        "学生学号",
        "学生姓名（可选）",
        "指导老师工号",
        "指导老师姓名（可选）",
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    # 示例数据
    examples = [
        ["S2021001", "张三", "T001", "李老师"],
        ["S2021002", "李四", "T001", "李老师"],
    ]
    
    for row_idx, example in enumerate(examples, 2):
        for col_idx, value in enumerate(example, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # 调整列宽
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 18
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=assignment_import_template.xlsx"},
    )
