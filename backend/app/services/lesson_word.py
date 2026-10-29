"""Create genuine, editable Word lesson plans, including tables/math/local images.

The model's text and historical lesson sources are parsed once on the server.
The browser and downloads both consume the resulting DOCX, never raw markup.
No URL or path from model output is fetched by the exporter.
"""
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree
from markdown_it import MarkdownIt
from PIL import Image
from .lesson_content import clean_lesson_content

WORD_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
MATH_TOKEN = re.compile(r'WORDFORMULA(\d+)TOKEN')
IMAGE_PATH = re.compile(r'^(?:/api/v1/textbooks/assets/|asset://)(tb-[a-f0-9]{32})(?:/image)?$')


def safe_filename(title):
    name = re.sub(r'[\x00-\x1f<>:"/\\|?*]', '_', title).strip(' .')[:100] or '教案'
    return name + ('' if name.endswith('教案') else '_教案') + '.docx'


def _math_element(name, *children):
    element = etree.Element(f'{{{MATH_NS}}}{name}', nsmap={'m': MATH_NS})
    for child in children:
        if child.tag==f'{{{MATH_NS}}}e' and name in ('e','num','den','sub','sup','lim','deg','oMath'):
            element.extend(list(child))
        else:
            element.append(child)
    return element


def _mathml_node(node):
    tag = etree.QName(node).localname
    children = [_mathml_node(child) for child in node]
    if tag in ('mi', 'mn', 'mo', 'mtext', 'ms'):
        run = _math_element('r')
        text = _math_element('t'); text.text = node.text or ''
        run.append(text)
        return run
    if tag in ('math', 'mrow', 'mstyle', 'mpadded', 'semantics'):
        return _math_element('e', *children)
    if tag == 'mfrac' and len(children) == 2:
        return _math_element('f', _math_element('num',children[0]), _math_element('den',children[1]))
    if tag in ('msup','msub','msubsup'):
        result = _math_element({'msup':'sSup','msub':'sSub','msubsup':'sSubSup'}[tag])
        result.append(_math_element('e',children[0]))
        for name,child in zip(('sub','sup') if tag=='msubsup' else ('sup',) if tag=='msup' else ('sub',),children[1:]):
            result.append(_math_element(name,child))
        return result
    if tag in ('msqrt','mroot'):
        props = _math_element('radPr')
        hidden = _math_element('degHide'); hidden.set(f'{{{MATH_NS}}}val','1' if tag=='msqrt' else '0'); props.append(hidden)
        degree = _math_element('deg', *([] if tag=='msqrt' else children[1:2]))
        return _math_element('rad',props,degree,_math_element('e',*(children if tag=='msqrt' else children[:1])))
    if tag in ('mover','munder','munderover'):
        base = children[0]
        if tag in ('munder','munderover'):
            base = _math_element('limLow',_math_element('e',base),_math_element('lim',children[1]))
        if tag in ('mover','munderover'):
            base = _math_element('limUpp',_math_element('e',base),_math_element('lim',children[-1]))
        return base
    if tag == 'mtable':
        return _math_element('m',*children)
    if tag in ('mtr','mlabeledtr'):
        return _math_element('mr',*children)
    if tag == 'mtd':
        return _math_element('e',*children)
    if tag == 'mfenced':
        props = _math_element('dPr')
        for name,value in [('begChr',node.get('open','(')),('endChr',node.get('close',')'))]:
            child = _math_element(name); child.set(f'{{{MATH_NS}}}val',value); props.append(child)
        return _math_element('d',props,_math_element('e',*children))
    if tag in ('annotation','annotation-xml'):
        return _math_element('e')
    return _math_element('e',*children)


def _formula(latex, display=False):
    xml = etree.fromstring(latex_to_mathml(latex).encode())
    converted = _mathml_node(xml)
    equation = _math_element('oMath', *list(converted))
    return _math_element('oMathPara',equation) if display else equation


def _extract_math(content):
    expressions = []
    def replace(match):
        text = match.group()
        display = text.startswith(('$$',r'\['))
        size = 2 if display or text.startswith(r'\(') else 1
        index = len(expressions)
        expressions.append((text[size:-size].strip(),display))
        return f'WORDFORMULA{index}TOKEN'
    pattern = r'\$\$[\s\S]+?\$\$|\\\[[\s\S]+?\\\]|\\\([\s\S]+?\\\)|(?<![\\$])\$(?!\$)[^$\n]+?\$'
    return re.sub(pattern,replace,content), expressions


def build_lesson_docx(content, title='教案', images=None):
    images = images or {}
    content = clean_lesson_content(content)
    wrapped = re.fullmatch(r'\s*```(?:markdown|md)?\s*\n([\s\S]*?)\n```\s*',content)
    if wrapped:
        content = wrapped.group(1)
    content, formulas = _extract_math(content)
    parser = MarkdownIt('commonmark',{'html':False}).enable('table').enable('strikethrough')
    tokens = parser.parse(content)
    document = Document()
    # python-docx templates can carry themed blue rules/colors. Remove them
    # explicitly so both Word and the browser honor the black title styles.
    for style in document.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        # docx-preview also follows linked character styles (TitleChar,
        # Heading1Char), which Word may override differently from the browser.
        if style.name.startswith(('Title','Heading','Subtitle')):
            style.font.color.rgb=RGBColor(0,0,0)
            style.font.underline=False
            for color in style.element.xpath('.//w:color'):
                for attr in ('themeColor','themeTint','themeShade'):
                    color.attrib.pop(qn('w:'+attr),None)
    section = document.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(2); section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.1); section.right_margin = Cm(2.1)
    usable = section.page_width-section.left_margin-section.right_margin
    for name,size,font in [('Normal',11,'宋体'),('Title',20,'黑体'),('Heading 1',15,'黑体'),('Heading 2',13,'黑体'),('Heading 3',12,'黑体'),('Caption',10,'宋体')]:
        style = document.styles[name]
        style.font.name = 'Times New Roman'; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0,0,0)
        style.font.underline = False
        rpr=style.element.get_or_add_rPr()
        fonts=rpr.get_or_add_rFonts(); fonts.set(qn('w:eastAsia'),font)
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            fonts.attrib.pop(qn('w:'+attr),None)
        for color in rpr.findall(qn('w:color')):
            for attr in ('themeColor','themeTint','themeShade'):
                color.attrib.pop(qn('w:'+attr),None)
        style.paragraph_format.line_spacing = 1.35
        style.paragraph_format.space_after = Pt(7 if name=='Normal' else 10)
        if name.startswith('Heading') or name=='Title':
            style.paragraph_format.keep_with_next = True
            style.font.bold = True
    title = re.sub(r'[#*`\n\r]', '', title).strip() or '教案'
    title_text = title if title.endswith('教案') else title+'教案'
    # Prefer the model's descriptive top heading when present, avoiding two titles.
    skip_heading = False
    if len(tokens)>1 and tokens[0].type=='heading_open' and tokens[0].tag=='h1':
        title_text = ''.join(t.content for t in tokens[1].children or [] if t.type in ('text','code_inline')) or title_text
        skip_heading = True
    title_paragraph = document.add_paragraph(title_text,'Title')
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.core_properties.title = title_text
    document.core_properties.subject = '教学设计'

    def append_text(paragraph,text,bold=False,italic=False,code=False):
        position = 0
        for match in MATH_TOKEN.finditer(text):
            if match.start()>position:
                run = paragraph.add_run(text[position:match.start()]); run.bold=bold;run.italic=italic
            latex,display = formulas[int(match.group(1))]
            try:
                paragraph._p.append(_formula(latex,display))
            except Exception:
                # Preserve failed formula source explicitly instead of dropping it.
                paragraph.add_run('公式待核对：'+latex)
            position = match.end()
        if position < len(text):
            run = paragraph.add_run(text[position:]);run.bold=bold;run.italic=italic
            if code:
                run.font.name='Consolas';run.font.size=Pt(10)

    def append_inline(paragraph,children,available_width=usable):
        bold,italic = False,False
        for token in children or []:
            if token.type=='strong_open': bold=True
            elif token.type=='strong_close': bold=False
            elif token.type=='em_open': italic=True
            elif token.type=='em_close': italic=False
            elif token.type in ('text','code_inline'):
                append_text(paragraph,token.content,bold,italic,token.type=='code_inline')
            elif token.type in ('softbreak','hardbreak'): paragraph.add_run().add_break()
            elif token.type=='image':
                match = IMAGE_PATH.fullmatch(token.attrGet('src') or '')
                path = images.get(match.group(1)) if match else None
                if path:
                    paragraph.paragraph_format.keep_together=True
                    with Image.open(path) as picture:
                        width,height = picture.size
                    # Fit tall illustrations as well as wide figures within A4.
                    image_width = min(available_width,int(Cm(11.5)),int(Cm(10)*width/max(height,1)))
                    picture=paragraph.add_run().add_picture(str(path),width=image_width)
                    picture._inline.docPr.set('descr',token.content or '教材图片')
                    if token.content:
                        caption=re.sub(r'^(图\s*[\d.\-－–]+)\s+\1\s*',r'\1 ',token.content)
                        paragraph.add_run('\n'+caption)
                else:
                    paragraph.add_run('图片待补充：'+(token.content or '教材配图'))
    
    def make_table(table_tokens):
        rows=[];row=None;cell=None
        for token in table_tokens:
            if token.type=='tr_open': row=[];rows.append(row)
            elif token.type in ('th_open','td_open'): cell=[];row.append(cell)
            elif token.type=='inline' and cell is not None: cell.extend(token.children or [])
        if not rows:return
        columns=max(len(row) for row in rows)
        table=document.add_table(rows=len(rows),cols=columns)
        table.autofit=False
        headers=[''.join(token.content for token in cell if token.type in ('text','code_inline')).strip() for cell in rows[0]]
        weights=[]
        for j in range(columns):
            header=headers[j] if j<len(headers) else ''
            if re.fullmatch(r'时间|时长|用时|分钟|序号|编号',header):weight=1
            elif re.search(r'环节|阶段|步骤|类别|项目|维度|要素',header):weight=.9 if columns==2 else 1.7
            else:weight=3.6
            weights.append(weight)
        ratios=[weight/sum(weights) for weight in weights]
        props=table._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for name in ('top','left','bottom','right','insideH','insideV'):
            border=OxmlElement('w:'+name);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'6');border.set(qn('w:color'),'D9D9D9');borders.append(border)
        props.append(borders)
        for col,ratio in zip(table.columns,ratios):col.width=int(usable*ratio)
        for i,row in enumerate(rows):
            for j,children in enumerate(row):
                target=table.cell(i,j);target.width=int(usable*ratios[j])
                tcpr=target._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                for name in ('top','left','bottom','right'):
                    margin=OxmlElement('w:'+name);margin.set(qn('w:w'),'100');margin.set(qn('w:type'),'dxa');margins.append(margin)
                tcpr.append(margins)
                align=OxmlElement('w:vAlign');align.set(qn('w:val'),'center');tcpr.append(align)
                paragraph=target.paragraphs[0];paragraph.paragraph_format.space_after=Pt(3);paragraph.paragraph_format.space_before=Pt(3)
                paragraph.paragraph_format.line_spacing=1.2
                append_inline(paragraph,children,max(int(Cm(1)),int(usable*ratios[j]-Cm(.4))))
                for run in paragraph.runs:run.font.size=Pt(10.5)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF5');tcpr.append(shade)
                    for run in paragraph.runs:run.bold=True
            if i==0:
                repeat=OxmlElement('w:tblHeader');table.rows[i]._tr.get_or_add_trPr().append(repeat)
        document.add_paragraph().paragraph_format.space_after=Pt(2)

    paragraph=None;lists=[];pending_item=False;skip_inline=False;index=0
    while index<len(tokens):
        token=tokens[index]
        if token.type=='table_open':
            end=next(i for i in range(index+1,len(tokens)) if tokens[i].type=='table_close')
            make_table(tokens[index:end+1]);index=end+1;continue
        if token.type=='heading_open':
            skip_inline=skip_heading and index==0
            paragraph=None if skip_inline else document.add_paragraph(style=f'Heading {min(int(token.tag[1:]),3)}')
        elif token.type=='paragraph_open':
            paragraph=document.add_paragraph(style='List Paragraph' if lists else 'Normal')
            if lists:
                paragraph.paragraph_format.left_indent=Pt(16*len(lists));paragraph.paragraph_format.first_line_indent=Pt(-12)
                if pending_item:
                    marker=str(lists[-1]['number'])+'. ' if lists[-1]['ordered'] else '• '
                    paragraph.add_run(marker);lists[-1]['number']+=1;pending_item=False
        elif token.type in ('bullet_list_open','ordered_list_open'):
            lists.append(dict(ordered=token.type=='ordered_list_open',number=int(token.attrGet('start') or 1)))
        elif token.type in ('bullet_list_close','ordered_list_close'):lists.pop()
        elif token.type=='list_item_open':pending_item=True
        elif token.type=='inline' and not skip_inline:
            if paragraph is None:paragraph=document.add_paragraph()
            append_inline(paragraph,token.children)
        elif token.type in ('paragraph_close','heading_close'):
            paragraph=None;skip_inline=False
        elif token.type in ('fence','code_block'):
            paragraph=document.add_paragraph()
            append_text(paragraph,token.content,code=True);paragraph=None
        index+=1
    # Keep short figure introductions on the same page as their illustration.
    paragraphs=document.paragraphs
    for previous,current in zip(paragraphs,paragraphs[1:]):
        if current._p.xpath('.//w:drawing') and 0<len(previous.text.strip())<=24:
            previous.paragraph_format.keep_with_next=True
    output=io.BytesIO();document.save(output)
    return output.getvalue()
