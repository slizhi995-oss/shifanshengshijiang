"""
生成"智教领航"系统试运营合作意向书 Word 文档
依赖：pip install python-docx
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

OUTPUT_PATH = "智教领航系统试运营合作意向书.docx"


def set_font(run, name="仿宋", size=12, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  font_name="仿宋", font_size=12, bold=False,
                  space_before=0, space_after=6, first_line_indent=True):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if first_line_indent:
        pf.first_line_indent = Pt(font_size * 2)
    run = p.add_run(text)
    set_font(run, name=font_name, size=font_size, bold=bold)
    return p


def add_heading(doc, text, level_name="一", font_size=14):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run(text)
    set_font(run, name="黑体", size=font_size, bold=True)
    return p


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run(text)
    set_font(run, name="黑体", size=22, bold=True)
    return p


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run(text)
    set_font(run, name="仿宋", size=12, bold=False, color=(128, 128, 128))
    return p


def add_blank(doc, lines=1):
    for _ in range(lines):
        p = doc.add_paragraph("")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)


def add_sign_line(doc, left_label, right_label):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    # 去除边框
    for cell in table.rows[0].cells:
        for border in ("top", "left", "bottom", "right"):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcBorders = OxmlElement("w:tcBorders")
            side = OxmlElement(f"w:{border}")
            side.set(qn("w:val"), "nil")
            tcBorders.append(side)
            tcPr.append(tcBorders)

    cells = table.rows[0].cells
    cells[0].width = Cm(8)
    cells[1].width = Cm(8)

    for i, label in enumerate([left_label, right_label]):
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(label)
        set_font(run, name="仿宋", size=12)
    return table


def main():
    doc = Document()

    # ── 页面边距
    for section in doc.sections:
        section.top_margin = Cm(3.0)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.2)
        section.right_margin = Cm(3.0)

    # ══════════════════════════════════════════
    # 标题
    # ══════════════════════════════════════════
    add_title(doc, "试运营合作意向书")
    add_subtitle(doc, '\u2014\u2014\u201c\u667a\u6559\u9886\u822a\uff1a\u57fa\u4e8eAI\u7684\u5e08\u8303\u751f\u8bd5\u8bb2\u5168\u6d41\u7a0b\u667a\u80fd\u8d4b\u80fd\u7cfb\u7edf\u201d')

    # 编号行
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run("编号：ZJ-COOP-2025-001")
    set_font(run, name="仿宋", size=10.5, color=(128, 128, 128))

    # ══════════════════════════════════════════
    # 甲乙双方信息框
    # ══════════════════════════════════════════
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    headers = ["甲  方（系统研发团队）", "乙  方（合作院校）"]
    contents = [
        "团队名称：智教领航教育科技团队\n所在学校：（请填写学校名称）\n联系人：（请填写负责人姓名）\n联系方式：（请填写联系方式）",
        "单位名称：西华师范大学数学科学学院\n地    址：四川省南充市顺庆区师大路1号\n联系人：（院长/联系人姓名）\n联系方式：（请填写联系方式）",
    ]
    for col_idx in range(2):
        hdr_cell = table.cell(0, col_idx)
        hdr_p = hdr_cell.paragraphs[0]
        hdr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr_p.paragraph_format.first_line_indent = Pt(0)
        hdr_run = hdr_p.add_run(headers[col_idx])
        set_font(hdr_run, name="黑体", size=12, bold=True)
        hdr_cell._tc.get_or_add_tcPr()

        cnt_cell = table.cell(1, col_idx)
        cnt_cell.paragraphs[0].clear()
        for line in contents[col_idx].split("\n"):
            cp = cnt_cell.add_paragraph(line)
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            cp.paragraph_format.first_line_indent = Pt(0)
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            for run in cp.runs:
                set_font(run, name="仿宋", size=11)

    add_blank(doc)

    # ══════════════════════════════════════════
    # 正文引言
    # ══════════════════════════════════════════
    add_paragraph(
        doc,
        "甲方（智教领航教育科技团队）与乙方（西华师范大学数学科学学院）经友好协商，就\u201c智教领航：基于AI的师范生试讲全流程智能赋能系统\u201d在乙方开展试运营合作事宜，本着自愿、平等、互利的原则，签订本合作意向书，以资共同遵守。",
    )

    # ══════════════════════════════════════════
    # 第一条
    # ══════════════════════════════════════════
    add_heading(doc, "第一条  合作背景与项目简介")
    add_paragraph(
        doc,
        "当前我国高等师范教育面临结构性矛盾：师范生在校人数已突破300万，而指导教师数量增长有限，一位导师通常需要指导25至35名学生。传统试讲评价以主观判断为主，缺乏统一量化标准，学生反馈等待周期长、备课缺乏互动训练，制约了师范生教学能力的培养质量。",
    )
    add_paragraph(
        doc,
        "针对上述问题，甲方研发了\u201c智教领航：基于AI的师范生试讲全流程智能赋能系统\u201d。系统集成通义千问系列三款专业大模型（文本模型qwen-plus、音频模型qwen-omni-turbo、视觉模型qwen-vl-max），围绕\u201c教案智能分析—视频多模态评估—AI对话式备课—AI辅助批改—学生成长画像\u201d五个环节构建完整训练闭环，支持管理员、指导教师、学生三端差异化使用，已完成全部核心功能开发并上线部署，覆盖33个功能页面、80余个API接口。",
    )
    add_paragraph(
        doc,
        "乙方作为西华师范大学数学科学学院，拥有丰富的师范生培养经验及充足的师范生用户群体，双方在学科背景、教育场景与技术能力上具有高度互补性，合作条件成熟。",
    )

    # ══════════════════════════════════════════
    # 第二条
    # ══════════════════════════════════════════
    add_heading(doc, "第二条  合作目的")
    add_paragraph(
        doc,
        "甲方希望借助乙方真实的师范生培养场景，对\u201c智教领航\u201d系统进行教学实践验证，收集真实用户使用数据，检验并优化系统功能与AI评估准确性；乙方希望通过引入智能化教学辅助工具，减轻指导教师批改工作量、提升师范生试讲训练质量，探索AI赋能师范教育的创新路径。",
    )

    # ══════════════════════════════════════════
    # 第三条
    # ══════════════════════════════════════════
    add_heading(doc, "第三条  合作内容与试运营安排")

    items = [
        ("（一）试运营范围", "乙方同意在数学科学学院师范生试讲课程教学中，选取不少于一个教学班（或批次）作为试点，组织师范生及指导教师使用\u201c智教领航\u201d系统完成至少一轮完整的试讲任务提交、AI分析与批改流程。"),
        ("（二）系统部署方式", "甲方为乙方提供在线访问账号或Docker私有化部署方案（双方协商确定），确保试运营期间系统稳定可用，乙方师生无需自行承担任何硬件或软件费用。"),
        ("（三）培训与支持", "甲方负责在试运营启动前为乙方管理员及指导教师提供系统操作培训（线上或线下均可），并在试运营期间提供技术支持与问题响应服务，响应时间不超过24小时。"),
        ("（四）数据与隐私", "试运营期间产生的学生教案文档、试讲视频及评分数据归乙方所有，甲方承诺仅用于系统功能优化与学术研究，不向任何第三方披露，并签署单独的数据保密协议。"),
        ("（五）试运营周期", "试运营周期为签署本意向书后______个月，双方可根据实际情况协商延续或转为正式合作。"),
    ]
    for title, content in items:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(title + "  ")
        set_font(r1, name="黑体", size=12, bold=True)
        r2 = p.add_run(content)
        set_font(r2, name="仿宋", size=12)

    # ══════════════════════════════════════════
    # 第四条
    # ══════════════════════════════════════════
    add_heading(doc, "第四条  双方权利与义务")

    add_paragraph(doc, "甲方的主要权利与义务：", first_line_indent=True)
    obligations_a = [
        "负责系统的技术开发、部署与维护，确保试运营期间系统稳定运行；",
        "为乙方提供系统使用培训及全程技术支持；",
        "在双方协商一致的前提下，可将本次合作作为典型案例用于学术交流和项目展示；",
        "尊重并保护乙方师生的个人信息与数据安全。",
    ]
    for i, item in enumerate(obligations_a, 1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        set_font(run, name="仿宋", size=12)

    add_paragraph(doc, "乙方的主要权利与义务：", first_line_indent=True)
    obligations_b = [
        "组织协调本院指导教师与师范生参与试运营，确保参与规模不低于双方约定的最低数量；",
        "积极收集师生使用反馈并定期向甲方反馈，协助甲方完善系统功能；",
        "在试运营期间不向院外第三方提供系统访问权限；",
        "配合甲方在学术交流场合对本次合作进行适度展示，包括但不限于竞赛、展览及论文发表。",
    ]
    for i, item in enumerate(obligations_b, 1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        set_font(run, name="仿宋", size=12)

    # ══════════════════════════════════════════
    # 第五条
    # ══════════════════════════════════════════
    add_heading(doc, "第五条  费用说明")
    add_paragraph(
        doc,
        "试运营期间，甲方向乙方免费提供系统使用权及全部技术支持服务，乙方无需支付任何费用。甲方承担系统服务器运营成本及AI模型API调用费用。双方如协商转为正式合作，将另行签订正式合作协议并约定具体费用标准。",
    )

    # ══════════════════════════════════════════
    # 第六条
    # ══════════════════════════════════════════
    add_heading(doc, "第六条  知识产权")
    add_paragraph(
        doc,
        "\u201c智教领航\u201d系统的软件著作权及全部核心技术成果归甲方所有（甲方已申请软件著作权《高校师范生试讲智慧管理系统》）。乙方在试运营期间对系统进行的二次配置（如Rubric评分标准个性化设置）的配置数据归乙方所有。双方合作期间共同产生的教学研究成果，如发表学术论文，应注明双方贡献，著作权归属另行协商。",
    )

    # ══════════════════════════════════════════
    # 第七条
    # ══════════════════════════════════════════
    add_heading(doc, "第七条  意向书性质与后续安排")
    add_paragraph(
        doc,
        "本意向书为双方合作意向的书面记录，不具有正式合同的法律约束力，仅代表双方就开展试运营合作的基本共识。双方应在本意向书签署后______日内，就试运营的具体安排进一步协商，并签订正式试运营合作协议。",
    )
    add_paragraph(
        doc,
        "本意向书一式两份，甲乙双方各执一份，自双方授权代表签字（盖章）之日起生效。",
    )

    # ══════════════════════════════════════════
    # 签字区
    # ══════════════════════════════════════════
    add_blank(doc, 2)

    sign_table = doc.add_table(rows=7, cols=2)
    sign_table.style = "Table Grid"
    # 去所有边框
    for row in sign_table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:val"), "nil")
                tcBorders.append(el)
            tcPr.append(tcBorders)

    sign_data = [
        ("甲  方（盖章）", "乙  方（盖章）"),
        ("智教领航教育科技团队", "西华师范大学数学科学学院"),
        ("", ""),
        ("授权代表签字：_______________", "授权代表签字（院长）：_______________"),
        ("", ""),
        ("日      期：______年____月____日", "日      期：______年____月____日"),
        ("", ""),
    ]
    for row_idx, (left, right) in enumerate(sign_data):
        cells = sign_table.rows[row_idx].cells
        for col_idx, text in enumerate([left, right]):
            p = cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx in (0, 1) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            bold = row_idx in (0, 1)
            font_name = "黑体" if row_idx in (0, 1) else "仿宋"
            set_font(run, name=font_name, size=12, bold=bold)

    # ══════════════════════════════════════════
    # 附件说明
    # ══════════════════════════════════════════
    add_blank(doc, 2)
    add_heading(doc, "附件：系统核心功能说明", font_size=12)

    annex_items = [
        ("附件1", "教案智能分析模块说明（七维度结构化分析、改进建议、AI评语生成）"),
        ("附件2", "视频多模态分析模块说明（语音分析+视觉教态分析指标体系）"),
        ("附件3", "AI备课助手功能说明（角色扮演式模拟课堂对话+练习报告）"),
        ("附件4", "学生成长画像模块说明（能力雷达图+成长趋势图+智能行动建议）"),
        ("附件5", "系统部署架构与数据安全说明（Docker私有化部署方案）"),
    ]
    for code, desc in annex_items:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"【{code}】  ")
        set_font(r1, name="黑体", size=11, bold=True)
        r2 = p.add_run(desc)
        set_font(r2, name="仿宋", size=11)

    doc.save(OUTPUT_PATH)
    print(f"✅ 文档已生成：{OUTPUT_PATH}")


if __name__ == "__main__":
    main()
